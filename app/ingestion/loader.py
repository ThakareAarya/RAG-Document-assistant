"""Document loading and chunking module for RAG ingestion pipeline."""
import os
from pathlib import Path
from typing import List, Dict, Any
from pydantic import BaseModel, Field

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    # Fallback to langchain.text_splitter if langchain_text_splitters isn't present
    from langchain.text_splitter import RecursiveCharacterTextSplitter

from app.config import settings


class Document(BaseModel):
    """Data model representing a document chunk with associated metadata."""
    page_content: str
    metadata: Dict[str, Any] = Field(default_factory=dict)


class DocumentLoader:
    """Parser supporting PDF, DOCX, TXT, and Markdown files."""

    @staticmethod
    def load_pdf(file_path: str | Path) -> List[Document]:
        """Extract text from a PDF file page by page using pypdf."""
        from pypdf import PdfReader
        
        file_path = Path(file_path)
        documents = []
        reader = PdfReader(str(file_path))
        
        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                documents.append(
                    Document(
                        page_content=text.strip(),
                        metadata={
                            "source": file_path.name,
                            "page": idx + 1,
                            "file_type": "pdf"
                        }
                    )
                )
        return documents

    @staticmethod
    def load_docx(file_path: str | Path) -> List[Document]:
        """Extract text from a DOCX file using python-docx."""
        import docx
        
        file_path = Path(file_path)
        doc = docx.Document(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        
        if not full_text:
            return []
            
        return [
            Document(
                page_content=full_text,
                metadata={
                    "source": file_path.name,
                    "page": 1,
                    "file_type": "docx"
                }
            )
        ]

    @staticmethod
    def load_text(file_path: str | Path) -> List[Document]:
        """Extract text from a TXT or MD file."""
        file_path = Path(file_path)
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
            
        if not text:
            return []
            
        return [
            Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "page": 1,
                    "file_type": file_path.suffix.lstrip(".").lower()
                }
            )
        ]

    @classmethod
    def load_file(cls, file_path: str | Path) -> List[Document]:
        """Auto-detect file extension and parse appropriately."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return cls.load_pdf(path)
        elif suffix in [".docx", ".doc"]:
            return cls.load_docx(path)
        elif suffix in [".txt", ".md", ".log", ".csv", ".json"]:
            return cls.load_text(path)
        else:
            # Fallback to plain text reader
            return cls.load_text(path)


def chunk_documents(
    documents: List[Document],
    chunk_size: int = None,
    chunk_overlap: int = None
) -> List[Document]:
    """Split documents into overlapping semantic chunks with preserved metadata.
    
    Args:
        documents: List of input raw Document objects.
        chunk_size: Maximum character count per chunk (default from settings).
        chunk_overlap: Overlap characters between consecutive chunks.
        
    Returns:
        List of split Document chunks with enriched metadata.
    """
    chunk_size = chunk_size or settings.CHUNK_SIZE
    chunk_overlap = chunk_overlap or settings.CHUNK_OVERLAP

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""]
    )

    chunked_docs: List[Document] = []
    chunk_counter = 0

    for doc in documents:
        splits = text_splitter.split_text(doc.page_content)
        for split_text in splits:
            if not split_text.strip():
                continue
            chunk_metadata = dict(doc.metadata)
            chunk_metadata.update({
                "chunk_id": chunk_counter,
                "char_count": len(split_text)
            })
            chunked_docs.append(
                Document(
                    page_content=split_text.strip(),
                    metadata=chunk_metadata
                )
            )
            chunk_counter += 1

    return chunked_docs
